#!/usr/bin/env python
"""Build Edit Tool–ready bankruptcy import xlsx from Athena join CSV.

Skips the usual manual steps:
  - zero-pads ZIP columns (leading zeros)
  - writes into the bankruptcy sample workbook shape
  - applies sample row formatting (dates d-mmm-yy, text @, etc.)

Example:
  python scripts\\build_bankruptcy_edit_tool_xlsx.py ^
    work\\2026-08-04\\bk_import_ready_20260731_20260803.csv ^
    --out work\\2026-08-04\\bk_edit_tool_import_20260731_20260803.xlsx
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection

DEFAULT_TEMPLATE = Path(r"C:\Users\<USERNAME>\Downloads\bankruptcy_sample (59).xlsx")

DATE_COLS = {
    "FILING_DATE",
    "RELEASE_DATE",
    "SCH_341_DATE",
    "DATE_RECORD_UPLOADED",
    "POCDATE",
}
ZIP_COLS = {"DEFENDANT_ZIP_CODE", "ATTORNEY_ZIP"}
# Keep these as text even when numeric-looking (IDs / codes that can lose shape).
FORCE_TEXT_COLS = ZIP_COLS | {
    "RMSID",
    "SAFE_NUMBER",
    "DEFENDANT_TYPE",
    "ALIAS_TYPE",
    "COURTID",
    "CASE_NUMBER",
    "JUDGES_INITIALS",
    "DEFENDANT_NAME",
    "DEFENDANT_NAME_GENERAT_CODE",
    "DEFENDANT_STREET_ADDRESS",
    "DEFENDANT_CITY",
    "DEFENDANT_STATE",
    "CASE_TYPE",
    "UNIQUE_CASE_ID",
    "ATTORNEY_ADDRESS",
    "ATTORNEY_CITY",
    "ATTORNEY_STATE",
    "ASSETS_AVAIL_FOR_DISTRIBUT",
    "SCH_341_TIME",
    "LAW_FIRM_ATTO_TRUSTEE_NAME",
    "ATTORNEY_NAME",
    "FILING_TYPES",  # can be like 7DNA / 11CO
}


def copy_cell_style(src, dest) -> None:
    if src.has_style:
        dest.font = Font(
            name=src.font.name,
            size=src.font.size,
            bold=src.font.bold,
            italic=src.font.italic,
            vertAlign=src.font.vertAlign,
            underline=src.font.underline,
            strike=src.font.strike,
            color=src.font.color,
        )
        dest.border = Border(
            left=src.border.left,
            right=src.border.right,
            top=src.border.top,
            bottom=src.border.bottom,
            diagonal=src.border.diagonal,
            diagonal_direction=src.border.diagonal_direction,
            outline=src.border.outline,
            vertical=src.border.vertical,
            horizontal=src.border.horizontal,
        )
        dest.fill = PatternFill(
            fill_type=src.fill.fill_type,
            start_color=src.fill.start_color,
            end_color=src.fill.end_color,
        )
        dest.number_format = src.number_format
        dest.protection = Protection(
            locked=src.protection.locked,
            hidden=src.protection.hidden,
        )
        dest.alignment = Alignment(
            horizontal=src.alignment.horizontal,
            vertical=src.alignment.vertical,
            textRotation=src.alignment.textRotation,
            wrap_text=src.alignment.wrap_text,
            shrink_to_fit=src.alignment.shrink_to_fit,
            indent=src.alignment.indent,
        )


def format_zip(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    if "-" in s:
        left, _, right = s.partition("-")
        if left.isdigit() and len(left) <= 5:
            return f"{left.zfill(5)}-{right}"
        return s
    if s.isdigit() and len(s) <= 5:
        return s.zfill(5)
    return s


def parse_date(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value if isinstance(value, datetime) else datetime(value.year, value.month, value.day)
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%m/%d/%Y", "%d-%b-%y", "%d-%b-%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    # Athena sometimes returns timestamps
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return s


def coerce_value(header: str, raw):
    if raw is None:
        return None
    s = str(raw).strip()
    if s == "":
        return None

    if header in ZIP_COLS:
        return format_zip(s)

    if header in DATE_COLS:
        return parse_date(s)

    if header in FORCE_TEXT_COLS:
        return s

    # numeric-ish optional fields
    if re.fullmatch(r"-?\d+", s):
        try:
            return int(s)
        except ValueError:
            return s
    if re.fullmatch(r"-?\d+\.\d+", s):
        try:
            return float(s)
        except ValueError:
            return s
    return s


def normalize_header(name: str) -> str:
    n = (name or "").strip()
    if n.lower() == "pocdate":
        return "POCDATE"
    return n


def load_rows(csv_path: Path) -> list[dict]:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = []
        for row in reader:
            rows.append({normalize_header(k): v for k, v in row.items()})
        return rows


def build_workbook(csv_path: Path, template_path: Path, out_path: Path) -> dict:
    rows = load_rows(csv_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, out_path)

    wb = load_workbook(out_path)
    ws = wb["bankruptcy"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    style_cells = [ws.cell(2, c) for c in range(1, ws.max_column + 1)]

    # clear sample data row
    for c in range(1, ws.max_column + 1):
        ws.cell(2, c).value = None

    zip_padded = 0
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, header in enumerate(headers, start=1):
            src_style = style_cells[c_idx - 1]
            cell = ws.cell(r_idx, c_idx)
            copy_cell_style(src_style, cell)

            if header in {None, "DELETED_BY_IMPORT", "DELETED_BY_EDITING_TOOL"}:
                cell.value = None
                continue

            raw = row.get(header)
            value = coerce_value(header, raw)
            if header in ZIP_COLS and value is not None:
                raw_s = str(raw).strip() if raw is not None else ""
                if raw_s.endswith(".0") and raw_s[:-2].isdigit():
                    raw_s = raw_s[:-2]
                if raw_s.isdigit() and len(raw_s) < 5:
                    zip_padded += 1
                cell.number_format = "@"
            cell.value = value

    wb.save(out_path)
    return {
        "rows": len(rows),
        "columns": len(headers),
        "zip_values_padded": zip_padded,
        "output": str(out_path),
        "template": str(template_path),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("csv", type=Path, help="Athena join CSV (bk_import_ready_*.csv)")
    parser.add_argument(
        "--template",
        type=Path,
        default=DEFAULT_TEMPLATE,
        help=f"Edit Tool sample workbook (default: {DEFAULT_TEMPLATE})",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="Output xlsx path (default: alongside CSV with _edit_tool.xlsx)",
    )
    args = parser.parse_args()

    if not args.csv.exists():
        raise SystemExit(f"CSV not found: {args.csv}")
    if not args.template.exists():
        raise SystemExit(f"Template not found: {args.template}")

    out = args.out or args.csv.with_name(args.csv.stem + "_edit_tool.xlsx")
    summary = build_workbook(args.csv, args.template, out)
    print(summary)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
