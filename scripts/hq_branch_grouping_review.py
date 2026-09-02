"""HQ/branch grouping review workbook. Review-only; no mutations."""

from __future__ import annotations

import argparse
import json
import os
import re
from datetime import date
from pathlib import Path

import oracledb
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
CLIENT = r"C:\Users\<USERNAME>\Documents\Codex\oracle-instant-client\instantclient_23_0"
DSN = "<ORACLE_HOST>:1521/<ORACLE_SERVICE>"
TARGET_DEFAULT = "<SAFE_017>"

# ponytail: one street-number + zip query per row; enrich SAFEs once at end


def connect():
    oracledb.init_oracle_client(lib_dir=CLIENT)
    c = oracledb.connect(user=os.environ["ORACLE_USER"], password=os.environ["ORACLE_PASSWORD"], dsn=os.environ.get("ORACLE_DSN", DSN))
    c.call_timeout = 90000
    return c


def q(conn, sql, params=None):
    cur = conn.cursor()
    cur.execute(sql, params or {})
    cols = [x[0].lower() for x in cur.description]
    return [{k: ("" if v is None else v) for k, v in zip(cols, r)} for r in cur.fetchall()]


def street_num(addr: str) -> str:
    m = re.search(r"(\d+)", addr or "")
    return m.group(1) if m else ""


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", (s or "").upper())).strip()


def addr_score(inp: dict, cand: dict) -> float:
    if inp["input_state"].upper() != str(cand.get("state", "")).upper():
        return 0.0
    s = 0.0
    if norm(inp["input_city"]) == norm(str(cand.get("city", ""))):
        s += 0.25
    iz = re.sub(r"\D", "", inp.get("input_zip", ""))[:5]
    cz = re.sub(r"\D", "", str(cand.get("zip", "")))[:5]
    if iz and cz == iz:
        s += 0.15
    it = set(norm(inp["input_address"]).split())
    ct = set(norm(str(cand.get("address", ""))).split())
    if it and ct:
        s += 0.6 * len(it & ct) / len(it | ct)
    return round(min(s, 1.0), 3)


def lvl(score: float) -> str:
    if score >= 0.95:
        return "EXACT"
    if score >= 0.8:
        return "VERY_SIMILAR"
    if score >= 0.55:
        return "SIMILAR"
    return "WEAK"


def classify(inp, cand, target, score, alvl):
    loc = str(cand.get("location_type", ""))
    hq = str(cand.get("hq_safe_number", "")).strip()
    safe = str(cand.get("safe_number", ""))
    cnt = int(cand.get("cnt_br") or 0)
    grp = str(cand.get("sn_in_grpstr", "N"))

    if safe == target and alvl in ("EXACT", "VERY_SIMILAR"):
        return "NO_ACTION", "HQ already at this address.", "N"
    if alvl == "WEAK":
        return "ADDRESS_MATCH_REVIEW", "Weak address match.", "Y"
    if loc == "2":
        if hq == target:
            return "ALREADY_BRANCH_UNDER_TARGET_HQ", "Already under target HQ.", "N"
        return "BRANCH_UNDER_DIFFERENT_HQ_REVIEW", f"Branch under {hq or 'unknown'}.", "Y"
    if loc == "1":
        if cnt > 0 or grp == "Y":
            return "SEPARATE_OR_COMPLEX_REVIEW", f"HQ complexity cnt_br={cnt} grp={grp}.", "Y"
        return "CONVERT_HQ_TO_BRANCH_REVIEW", "Standalone HQ at branch address.", "Y"
    return "ADDRESS_MATCH_REVIEW", f"Unexpected loc={loc}.", "Y"


def search_row(inp: dict, target: str) -> list[dict]:
    num = street_num(inp["input_address"])
    if not num:
        return []
    zip5 = re.sub(r"\D", "", inp.get("input_zip", ""))[:5]
    sql = """
    SELECT * FROM (
      SELECT a.safe_number, a.company_name, a.address, a.city, a.state, a.zip,
             a.location_type, a.hq_safe_number
      FROM company_search_view a
      WHERE a.state = :st AND a.zip = :zip AND a.address LIKE :pfx
      ORDER BY CASE WHEN a.hq_safe_number = :hq THEN 0 ELSE 1 END, a.location_type
    ) WHERE ROWNUM <= 20
    """
    with connect() as conn:
        return q(conn, sql, {"st": inp["input_state"].upper(), "zip": zip5, "pfx": f"{num}%", "hq": target})


def enrich(target: str, safes: list[str]) -> dict[str, dict]:
    if not safes:
        return {}
    ph = ", ".join(f"'{s}'" for s in sorted(set(safes)))
    sql = f"""
    SELECT reporting_utils.fn_br_count_for_hq(a.safe_number) cnt_br,
           CASE WHEN a.safe_number IN (SELECT safe_number FROM core_data.group_structure WHERE safe_number IS NOT NULL)
                THEN 'Y' ELSE 'N' END sn_in_grpstr,
           a.safe_number, a.company_name, a.address, a.city, a.state, a.zip,
           a.location_type, a.hq_safe_number
    FROM company_search_view a WHERE a.safe_number IN ({ph})
    """
    with connect() as conn:
        rows = q(conn, sql)
    return {str(r["safe_number"]): r for r in rows}


def hq_context(target: str) -> tuple[dict, list[dict]]:
    base = """
    SELECT reporting_utils.fn_br_count_for_hq(a.safe_number) cnt_br,
           CASE WHEN a.safe_number IN (SELECT safe_number FROM core_data.group_structure WHERE safe_number IS NOT NULL)
                THEN 'Y' ELSE 'N' END sn_in_grpstr,
           a.safe_number, a.company_name, a.address, a.city, a.state, a.zip,
           a.location_type, a.hq_safe_number
    FROM company_search_view a WHERE {where}
    """
    with connect() as conn:
        hq = q(conn, base.format(where="a.safe_number = :s"), {"s": target})
        branches = q(conn, base.format(where="a.hq_safe_number = :s"), {"s": target})
    return (hq[0] if hq else {}), branches


def sheet(wb, name, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    if first:
        ws.title = name
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="18324A")
    for r in rows:
        ws.append([r.get(h, "") for h in headers])


def run(target: str, inputs: list[dict], out_dir: Path) -> dict:
    hq, branches = hq_context(target)
    hq_name = str(hq.get("company_name", ""))

    raw_matches: list[tuple[dict, list[dict]]] = []
    all_safes = [target]
    for inp in inputs:
        print(f"seq {inp['input_seq']} {inp['input_city']}", flush=True)
        raw_matches.append((inp, search_row(inp, target)))
        all_safes.extend(str(m["safe_number"]) for m in raw_matches[-1][1])

    enriched = enrich(target, all_safes)

    input_rows, cand_rows, action_rows, create_rows, review_rows = [], [], [], [], []

    for inp, matches in raw_matches:
        input_rows.append({
            "input_seq": inp["input_seq"], "target_hq_safe_number": target, "target_hq_name": hq_name,
            "input_branch_name": inp.get("input_branch_name", ""), "input_address": inp["input_address"],
            "input_city": inp["input_city"], "input_state": inp["input_state"], "input_zip": inp["input_zip"],
            "notes": inp.get("notes", ""),
        })
        scored = []
        for m in matches:
            c = {**m, **enriched.get(str(m["safe_number"]), {})}
            sc = addr_score(inp, c)
            al = lvl(sc)
            nm = "EXACT" if norm(inp.get("input_branch_name", hq_name)) in norm(str(c.get("company_name", ""))) else "UNKNOWN"
            cand_rows.append({
                "input_seq": inp["input_seq"], "input_address": inp["input_address"],
                "candidate_safe_number": c.get("safe_number"), "candidate_company_name": c.get("company_name"),
                "candidate_address": c.get("address"), "candidate_city": c.get("city"),
                "candidate_state": c.get("state"), "candidate_zip": c.get("zip"),
                "location_type": c.get("location_type"), "current_hq_safe_number": c.get("hq_safe_number"),
                "cnt_br": c.get("cnt_br"), "sn_in_grpstr": c.get("sn_in_grpstr"),
                "address_match_level": al, "name_match_level": nm,
                "confidence": "High" if al == "EXACT" else "Medium" if al == "VERY_SIMILAR" else "Low",
                "notes": f"score={sc}",
            })
            scored.append((c, sc, al))

        if not scored:
            term = f"state={inp['input_state']}; zip={inp['input_zip']}; num={street_num(inp['input_address'])}%"
            create_rows.append({
                "input_seq": inp["input_seq"], "branch_name_if_available": inp.get("input_branch_name", ""),
                "input_address": inp["input_address"], "input_city": inp["input_city"],
                "input_state": inp["input_state"], "input_zip": inp["input_zip"],
                "why_possible_create": "No Oracle row at street number + ZIP.",
                "searched_terms": term, "notes": "",
            })
            action_rows.append({
                "input_seq": inp["input_seq"], "input_address": inp["input_address"],
                "candidate_safe_number": "", "candidate_company_name": "", "current_location_type": "",
                "current_hq_safe_number": "", "target_hq_safe_number": target,
                "action_flag": "POSSIBLE_CREATE", "recommended_action": "Review Company Import create.",
                "reason": "No address match.", "manual_review_needed": "Y", "notes": term,
            })
            continue

        best = max(scored, key=lambda x: x[1])
        c, sc, al = best
        flag, rec, manual = classify(inp, c, target, sc, al)
        action_rows.append({
            "input_seq": inp["input_seq"], "input_address": inp["input_address"],
            "candidate_safe_number": c.get("safe_number"), "candidate_company_name": c.get("company_name"),
            "current_location_type": c.get("location_type"), "current_hq_safe_number": c.get("hq_safe_number"),
            "target_hq_safe_number": target, "action_flag": flag, "recommended_action": rec,
            "reason": rec, "manual_review_needed": manual, "notes": f"score={sc}",
        })
        if len(scored) > 1:
            review_rows.append({
                "input_seq": inp["input_seq"], "input_address": inp["input_address"],
                "possible_safe_numbers": ", ".join(str(x[0].get("safe_number")) for x in scored[:5]),
                "issue_type": "MULTIPLE_CANDIDATES", "why_review_only": f"{len(scored)} hits; best score={sc}",
                "notes": "",
            })

    branch_rows = [{
        "target_hq_safe_number": target, "branch_safe_number": b.get("safe_number"),
        "branch_company_name": b.get("company_name"), "branch_address": b.get("address"),
        "branch_city": b.get("city"), "branch_state": b.get("state"), "branch_zip": b.get("zip"),
        "location_type": b.get("location_type"), "cnt_br": b.get("cnt_br"), "sn_in_grpstr": b.get("sn_in_grpstr"), "notes": "",
    } for b in branches]

    wb = Workbook()
    sheet(wb, "input_addresses",
          ["input_seq", "target_hq_safe_number", "target_hq_name", "input_branch_name", "input_address", "input_city", "input_state", "input_zip", "notes"],
          input_rows, True)
    sheet(wb, "current_hq_branches",
          ["target_hq_safe_number", "branch_safe_number", "branch_company_name", "branch_address", "branch_city", "branch_state", "branch_zip", "location_type", "cnt_br", "sn_in_grpstr", "notes"],
          branch_rows)
    sheet(wb, "candidate_address_matches",
          ["input_seq", "input_address", "candidate_safe_number", "candidate_company_name", "candidate_address", "candidate_city", "candidate_state", "candidate_zip", "location_type", "current_hq_safe_number", "cnt_br", "sn_in_grpstr", "address_match_level", "name_match_level", "confidence", "notes"],
          cand_rows)
    sheet(wb, "recommended_actions",
          ["input_seq", "input_address", "candidate_safe_number", "candidate_company_name", "current_location_type", "current_hq_safe_number", "target_hq_safe_number", "action_flag", "recommended_action", "reason", "manual_review_needed", "notes"],
          action_rows)
    sheet(wb, "possible_creates",
          ["input_seq", "branch_name_if_available", "input_address", "input_city", "input_state", "input_zip", "why_possible_create", "searched_terms", "notes"],
          create_rows)
    sheet(wb, "review_only",
          ["input_seq", "input_address", "possible_safe_numbers", "issue_type", "why_review_only", "notes"],
          review_rows)

    out = out_dir / "hq_branch_grouping_review.xlsx"
    out_dir.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    counts = {
        "target_hq_safe_number": target, "target_hq_name": hq_name,
        "input_addresses": len(inputs), "current_branches": len(branches),
        "already_correct": sum(1 for r in action_rows if r["action_flag"] == "ALREADY_BRANCH_UNDER_TARGET_HQ"),
        "no_action": sum(1 for r in action_rows if r["action_flag"] == "NO_ACTION"),
        "candidates_found": len(cand_rows),
        "convert_hq_review": sum(1 for r in action_rows if r["action_flag"] == "CONVERT_HQ_TO_BRANCH_REVIEW"),
        "different_hq_review": sum(1 for r in action_rows if r["action_flag"] == "BRANCH_UNDER_DIFFERENT_HQ_REVIEW"),
        "complex_review": sum(1 for r in action_rows if r["action_flag"] == "SEPARATE_OR_COMPLEX_REVIEW"),
        "possible_creates": len(create_rows),
        "review_only": len(review_rows),
        "workbook_path": str(out),
    }
    print(json.dumps(counts, indent=2))
    return counts


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--target-hq", default=TARGET_DEFAULT)
    p.add_argument("--addresses-json", type=Path, required=True)
    p.add_argument("--output-dir", type=Path)
    a = p.parse_args()
    inputs = json.loads(a.addresses_json.read_text(encoding="utf-8"))
    run(a.target_hq.upper(), inputs, a.output_dir or ROOT / "work" / date.today().isoformat())


if __name__ == "__main__":
    main()
