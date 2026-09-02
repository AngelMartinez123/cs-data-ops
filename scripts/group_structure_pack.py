"""Build group-structure action workbook + pre/post-check SQL from a candidate list."""

from __future__ import annotations

import argparse
import csv
import textwrap
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent

REVIEW_HEADERS = [
    "source_safe_number",
    "target_uhc_safe_number",
    "action_type",
    "company_id",
    "company_name",
    "branch_count",
    "sn_in_grpstr",
    "confidence",
    "evidence",
    "caution_notes",
    "action_status",
]

ACTION_TYPES = {"move", "add", "branch_separate", "review"}


def precheck_sql(safe_numbers: list[str]) -> str:
    in_list = ",\n    ".join(f"'{s}'" for s in safe_numbers)
    return textwrap.dedent(
        f"""
        -- Pre-check: branch count + group_structure presence
        SELECT
            reporting_utils.fn_br_count_for_hq(a.safe_number) AS cnt_br,
            CASE
                WHEN a.safe_number IN (
                    SELECT safe_number
                    FROM core_data.group_structure
                    WHERE safe_number IS NOT NULL
                ) THEN 'Y'
                ELSE 'N'
            END AS sn_in_grpstr,
            a.*
        FROM company_search_view a
        WHERE a.safe_number IN (
            {in_list}
        );
        """
    ).strip()


def postcheck_sql(safe_numbers: list[str]) -> str:
    return precheck_sql(safe_numbers).replace(
        "-- Pre-check:", "-- Post-check (re-run after actions):"
    )


def move_sql(source_safe: str, target_uhc: str, company_id: str) -> str:
    cid = company_id.strip() if company_id else "<RESOLVE_COMPANY_ID>"
    return textwrap.dedent(
        f"""
        -- MOVE: {source_safe} -> target UHC {target_uhc}
        BEGIN
            editing_api.pkg_entity_group.move_to_new_safe_number(
                i_safe_number => '{target_uhc}',
                i_companyid   => {cid}
            );
            COMMIT;
        EXCEPTION
            WHEN OTHERS THEN
                ROLLBACK;
                RAISE;
        END;
        /
        """
    ).strip()


def add_sql(source_safe: str, target_uhc: str) -> str:
    return textwrap.dedent(
        f"""
        -- ADD subsidiary: {source_safe} under UHC {target_uhc}
        -- Use PKG_ENTITY_RELATIONS (not GROUP — avoids full group rebuild)
        SET SERVEROUTPUT ON;
        DECLARE
            l_parent_companyid VARCHAR2(25);
            l_cur SYS_REFCURSOR;
        BEGIN
            l_parent_companyid := reporting_utils.fn_companyid('{target_uhc}');
            editing_api.pkg_entity_relations.set_company_relation(
                i_parentid        => l_parent_companyid,
                i_companyid       => NULL,
                i_safe_number     => '{source_safe}',
                i_company_name    => NULL,
                i_company_type    => 'SUBSIDIARY',
                i_company_address => NULL,
                i_company_city    => NULL,
                i_company_state   => NULL,
                i_company_zipcode => NULL,
                i_company_phone   => NULL,
                i_company_country => 'UNITED STATES OF AMERICA',
                i_company_number  => NULL,
                i_expiry_date     => ADD_MONTHS(SYSDATE, 1200),
                i_dedupe_override => 'N',
                i_is_hidden       => 'N',
                o_cursor          => l_cur
            );
            IF l_cur%ISOPEN THEN
                CLOSE l_cur;
            END IF;
            COMMIT;
        EXCEPTION
            WHEN OTHERS THEN
                IF l_cur%ISOPEN THEN
                    CLOSE l_cur;
                END IF;
                ROLLBACK;
                RAISE;
        END;
        /
        """
    ).strip()


def branch_separate_sql(branch_safe: str) -> str:
    return textwrap.dedent(
        f"""
        -- BRANCH SEPARATION: hide branch relationship, then promote to HQ (loc type 1)
        -- Step 1: hide active branch relationship (fill row fields from core_data.company_relationship)
        SET SERVEROUTPUT ON;
        DECLARE
            p_branch_safe_number VARCHAR2(25) := '{branch_safe}';
            r core_data.company_relationship%ROWTYPE;
        BEGIN
            SELECT *
            INTO r
            FROM core_data.company_relationship
            WHERE cs_company_id = p_branch_safe_number
              AND company_type = 'BRANCH'
              AND deleted_by_import IS NULL
              AND deleted_by_editing_tool IS NULL
              AND ROWNUM = 1;

            editing_api.pkg_entity_relations.set_company_relation(
                i_parentid        => r.parentid,
                i_companyid       => r.companyid,
                i_safe_number     => r.cs_company_id,
                i_company_name    => r.company_name,
                i_company_type    => r.company_type,
                i_company_address => r.company_address,
                i_company_city    => r.company_city,
                i_company_state   => r.company_state,
                i_company_zip     => r.company_zip,
                i_company_country => r.company_country,
                i_is_ultimate     => 'N',
                i_expiry_date     => SYSDATE,
                i_reason          => 'Branch separation - hide relationship'
            );
            COMMIT;
        END;
        /

        -- Step 2: promote company data to location_type = 1 with future expiry (adjust dates as needed)
        -- Use your approved promotion procedure / Edit Tool SQL block here for {branch_safe}.
        """
    ).strip()


def load_candidates(path: Path, default_target: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        for raw in reader:
            action = (raw.get("action_type") or "review").strip().lower()
            if action not in ACTION_TYPES:
                raise ValueError(f"Unknown action_type '{action}' for {raw}")
            rows.append(
                {
                    "source_safe_number": (raw.get("source_safe_number") or "").strip(),
                    "target_uhc_safe_number": (
                        raw.get("target_uhc_safe_number") or default_target
                    ).strip(),
                    "action_type": action,
                    "company_id": (raw.get("company_id") or "").strip(),
                    "company_name": (raw.get("company_name") or "").strip(),
                    "confidence": (raw.get("confidence") or "").strip(),
                    "evidence": (raw.get("evidence") or "").strip(),
                    "caution_notes": (raw.get("caution_notes") or raw.get("notes") or "").strip(),
                    "branch_count": (raw.get("branch_count") or "").strip(),
                    "sn_in_grpstr": (raw.get("sn_in_grpstr") or "").strip(),
                    "action_status": (raw.get("action_status") or "pending").strip(),
                }
            )
    if not rows:
        raise ValueError(f"No rows in {path}")
    return rows


def write_sheet(ws, headers: list[str], data_rows: list[list[str]]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in data_rows:
        ws.append(row)
    ws.freeze_panes = "A2"


def build_workbook(
    candidates: list[dict[str, str]],
    output_xlsx: Path,
    precheck_path: Path,
    postcheck_path: Path,
) -> None:
    safe_numbers = sorted(
        {c["source_safe_number"] for c in candidates if c["source_safe_number"]}
        | {c["target_uhc_safe_number"] for c in candidates if c["target_uhc_safe_number"]}
    )
    precheck_path.write_text(precheck_sql(safe_numbers) + "\n", encoding="utf-8")
    postcheck_path.write_text(postcheck_sql(safe_numbers) + "\n", encoding="utf-8")

    review_rows = [
        [
            c["source_safe_number"],
            c["target_uhc_safe_number"],
            c["action_type"],
            c["company_id"],
            c["company_name"],
            c["branch_count"],
            c["sn_in_grpstr"],
            c["confidence"],
            c["evidence"],
            c["caution_notes"],
            c["action_status"],
        ]
        for c in candidates
    ]

    move_rows: list[list[str]] = []
    add_rows: list[list[str]] = []
    review_only: list[list[str]] = []
    oob_rows: list[list[str]] = []

    for c in candidates:
        src = c["source_safe_number"]
        tgt = c["target_uhc_safe_number"]
        action = c["action_type"]
        if action == "move":
            move_rows.append([src, tgt, move_sql(src, tgt, c["company_id"])])
        elif action == "add":
            add_rows.append([src, tgt, add_sql(src, tgt)])
        elif action == "branch_separate":
            move_rows.append([src, tgt, branch_separate_sql(src)])
        elif action == "review":
            review_only.append(
                [src, tgt, c["company_name"], c["evidence"], c["caution_notes"]]
            )

    wb = Workbook()
    ws_review = wb.active
    ws_review.title = "group_structure_review"
    write_sheet(ws_review, REVIEW_HEADERS, review_rows)

    ws_oob = wb.create_sheet("oob_closed_action")
    write_sheet(ws_oob, ["source_safe_number", "notes", "evidence", "action_status"], oob_rows)

    ws_move = wb.create_sheet("move_to_new_group_sql")
    write_sheet(
        ws_move,
        ["source_safe_number", "target_uhc_safe_number", "sql_block"],
        move_rows,
    )

    ws_add = wb.create_sheet("add_company_sql")
    write_sheet(
        ws_add,
        ["source_safe_number", "target_uhc_safe_number", "sql_block"],
        add_rows,
    )

    ws_not = wb.create_sheet("review_not_actioned")
    write_sheet(
        ws_not,
        ["source_safe_number", "target_uhc_safe_number", "company_name", "evidence", "notes"],
        review_only,
    )

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate group structure action pack.")
    parser.add_argument(
        "--input",
        "-i",
        type=Path,
        required=True,
        help="CSV with source_safe_number, action_type, target_uhc_safe_number, ...",
    )
    parser.add_argument(
        "--target-uhc",
        required=True,
        help="Default target UHC SAFE when row omits target_uhc_safe_number.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Output directory (default: work/YYYY-MM-DD).",
    )
    parser.add_argument(
        "--prefix",
        default="group_structure_pack",
        help="Output file prefix.",
    )
    args = parser.parse_args()

    out_dir = args.output_dir or (ROOT / "work" / date.today().isoformat())
    out_dir.mkdir(parents=True, exist_ok=True)

    candidates = load_candidates(args.input, args.target_uhc)
    xlsx = out_dir / f"{args.prefix}.xlsx"
    pre = out_dir / f"{args.prefix}_precheck.sql"
    post = out_dir / f"{args.prefix}_postcheck.sql"

    build_workbook(candidates, xlsx, pre, post)
    print(f"Wrote {xlsx}")
    print(f"Wrote {pre}")
    print(f"Wrote {post}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
