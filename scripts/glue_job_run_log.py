#!/usr/bin/env python3
"""Fetch AWS Glue job run history and write a pipeline timing workbook."""

from __future__ import annotations

import argparse
import statistics
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import boto3
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ET = ZoneInfo("America/New_York")
UTC = timezone.utc

MABEL_JOB = "us_mabel_verified_v001"
BK_JOB = "bankruptcy-unmatched-daily"
BK_MIN_EXPORT_SECONDS = 90


def fetch_job_runs(job_name: str, profile: str, region: str, since: datetime) -> list[dict]:
    session = boto3.Session(profile_name=profile, region_name=region)
    glue = session.client("glue")
    all_runs: list[dict] = []
    token = None
    while True:
        kwargs = {"JobName": job_name, "MaxResults": 200}
        if token:
            kwargs["NextToken"] = token
        resp = glue.get_job_runs(**kwargs)
        batch = resp.get("JobRuns", [])
        all_runs.extend(batch)
        token = resp.get("NextToken")
        if not token:
            break
        oldest = batch[-1].get("StartedOn") if batch else None
        if oldest and oldest.astimezone(UTC) < since:
            break
    return [
        r
        for r in all_runs
        if r.get("StartedOn") and r["StartedOn"].astimezone(UTC) >= since
    ]


def duration_seconds(started, completed) -> float | None:
    if not started or not completed:
        return None
    return (completed - started).total_seconds()


def fmt_time(dt: datetime | None) -> str:
    if dt is None:
        return ""
    return dt.astimezone(ET).strftime("%H:%M:%S")


def fmt_datetime(dt: datetime | None) -> str:
    if dt is None:
        return ""
    return dt.astimezone(ET).strftime("%Y-%m-%d %H:%M:%S")


def seconds_since_midnight(dt: datetime) -> int:
    local = dt.astimezone(ET)
    return local.hour * 3600 + local.minute * 60 + local.second


def seconds_to_time_str(total_seconds: float) -> str:
    total_seconds = int(round(total_seconds)) % (24 * 3600)
    h = total_seconds // 3600
    m = (total_seconds % 3600) // 60
    s = total_seconds % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def minutes_label(minutes: float | None) -> str:
    if minutes is None:
        return ""
    sign = "+" if minutes >= 0 else "-"
    return f"{sign}{abs(minutes):.1f} min"


def weekday_name(dt: datetime) -> str:
    return dt.astimezone(ET).strftime("%A")


def day_key(dt: datetime) -> str:
    return dt.astimezone(ET).date().isoformat()


def autosize_columns(ws):
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 44)


def style_header_row(ws, row: int = 1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def append_run_log_sheet(wb: Workbook, sheet_name: str, job_name: str, runs: list[dict]):
    ws = wb.create_sheet(sheet_name)
    headers = [
        "run_date_et",
        "weekday",
        "started_et",
        "completed_et",
        "duration_min",
        "duration_sec",
        "started_utc",
        "completed_utc",
        "job_run_state",
        "execution_time_sec",
        "job_run_id",
        "error_message",
    ]
    ws.append(headers)
    for r in sorted(runs, key=lambda item: item["StartedOn"]):
        started = r["StartedOn"]
        completed = r.get("CompletedOn")
        dur = duration_seconds(started, completed)
        ws.append(
            [
                day_key(started),
                weekday_name(started),
                fmt_time(started),
                fmt_time(completed),
                round(dur / 60, 2) if dur is not None else "",
                round(dur, 1) if dur is not None else "",
                started.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S"),
                completed.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S") if completed else "",
                r.get("JobRunState", ""),
                r.get("ExecutionTime"),
                r.get("Id", ""),
                r.get("ErrorMessage", ""),
            ]
        )
    style_header_row(ws)
    autosize_columns(ws)
    return ws


def summarize_times(runs: list[dict], label: str) -> list[tuple[str, str, str]]:
    succeeded = [r for r in runs if r.get("JobRunState") == "SUCCEEDED" and r.get("StartedOn")]
    if not succeeded:
        return [(f"{label}_runs", "0", "No runs in period")]

    start_seconds = [seconds_since_midnight(r["StartedOn"]) for r in succeeded]
    complete_seconds = [
        seconds_since_midnight(r["CompletedOn"])
        for r in succeeded
        if r.get("CompletedOn")
    ]
    durations = [
        duration_seconds(r["StartedOn"], r.get("CompletedOn"))
        for r in succeeded
        if r.get("CompletedOn")
    ]
    rows = [
        (f"{label}_total_runs", str(len(runs)), "All job runs in period"),
        (f"{label}_successful_runs", str(len(succeeded)), "Succeeded runs only"),
        (f"{label}_earliest_start", seconds_to_time_str(min(start_seconds)), "Earliest StartedOn (ET)"),
        (f"{label}_latest_start", seconds_to_time_str(max(start_seconds)), "Latest StartedOn (ET)"),
        (f"{label}_average_start", seconds_to_time_str(statistics.mean(start_seconds)), "Mean start time of day (ET)"),
        (f"{label}_median_start", seconds_to_time_str(statistics.median(start_seconds)), "Median start time of day (ET)"),
    ]
    if complete_seconds:
        rows.extend(
            [
                (f"{label}_earliest_complete", seconds_to_time_str(min(complete_seconds)), "Earliest CompletedOn (ET)"),
                (f"{label}_latest_complete", seconds_to_time_str(max(complete_seconds)), "Latest CompletedOn (ET)"),
                (
                    f"{label}_average_complete",
                    seconds_to_time_str(statistics.mean(complete_seconds)),
                    "Mean completion time of day (ET)",
                ),
                (
                    f"{label}_median_complete",
                    seconds_to_time_str(statistics.median(complete_seconds)),
                    "Median completion time of day (ET)",
                ),
            ]
        )
    if durations:
        rows.extend(
            [
                (f"{label}_average_duration_min", str(round(statistics.mean(durations) / 60, 2)), "Mean wall-clock runtime"),
                (f"{label}_median_duration_min", str(round(statistics.median(durations) / 60, 2)), "Median wall-clock runtime"),
            ]
        )
    return rows


def first_run_by_day(runs: list[dict]) -> dict[str, dict]:
    grouped: dict[str, list[dict]] = {}
    for run in runs:
        grouped.setdefault(day_key(run["StartedOn"]), []).append(run)
    return {day: sorted(day_runs, key=lambda item: item["StartedOn"])[0] for day, day_runs in grouped.items()}


def export_run_by_day(runs: list[dict], min_seconds: float = BK_MIN_EXPORT_SECONDS) -> dict[str, dict | None]:
    grouped: dict[str, list[dict]] = {}
    for run in runs:
        if run.get("JobRunState") != "SUCCEEDED":
            continue
        grouped.setdefault(day_key(run["StartedOn"]), []).append(run)

    exports: dict[str, dict | None] = {}
    for day, day_runs in grouped.items():
        best = max(day_runs, key=lambda item: duration_seconds(item["StartedOn"], item.get("CompletedOn")) or 0)
        dur = duration_seconds(best["StartedOn"], best.get("CompletedOn")) or 0
        exports[day] = best if dur >= min_seconds else None
    return exports


def build_workbook(
    mabel_runs: list[dict],
    bk_runs: list[dict],
    since_label: str,
    output_path: Path,
):
    wb = Workbook()
    wb.remove(wb.active)

    append_run_log_sheet(wb, "mabel_run_log", MABEL_JOB, mabel_runs)
    append_run_log_sheet(wb, "bk_run_log", BK_JOB, bk_runs)

    mabel_first = first_run_by_day([r for r in mabel_runs if r.get("JobRunState") == "SUCCEEDED"])
    bk_exports = export_run_by_day(bk_runs)

    ws_pipe = wb.create_sheet("daily_pipeline")
    ws_pipe.append(
        [
            "run_date_et",
            "weekday",
            "mabel_start_et",
            "mabel_complete_et",
            "mabel_duration_min",
            "bk_matching_start_et",
            "bk_matching_complete_et",
            "bk_export_duration_min",
            "wait_mabel_to_bk_start_min",
            "total_mabel_start_to_bk_complete_min",
            "bk_poll_attempts_that_day",
            "delay_vs_0800_bk_complete",
            "notes",
        ]
    )
    style_header_row(ws_pipe)

    pipeline_rows = []
    target_8am = 8 * 3600
    bk_by_day: dict[str, list[dict]] = {}
    for run in bk_runs:
        bk_by_day.setdefault(day_key(run["StartedOn"]), []).append(run)

    for day in sorted(set(mabel_first) | set(bk_exports)):
        mabel = mabel_first.get(day)
        export = bk_exports.get(day)
        notes = []

        if not mabel:
            notes.append("No mabel run")
        if export is None and day in bk_by_day:
            notes.append("No BK export detected (max run < 90s)")
        elif export is None:
            notes.append("No BK runs")

        mabel_start = mabel["StartedOn"] if mabel else None
        mabel_complete = mabel.get("CompletedOn") if mabel else None
        mabel_dur = duration_seconds(mabel_start, mabel_complete) if mabel else None

        bk_start = export["StartedOn"] if export else None
        bk_complete = export.get("CompletedOn") if export else None
        bk_dur = duration_seconds(bk_start, bk_complete) if export else None

        wait_min = None
        total_min = None
        if mabel_complete and bk_start:
            wait_min = (bk_start - mabel_complete).total_seconds() / 60
        if mabel_start and bk_complete:
            total_min = (bk_complete - mabel_start).total_seconds() / 60

        delay_label = ""
        if bk_complete:
            delay = seconds_since_midnight(bk_complete) - target_8am
            delay_label = minutes_label(delay / 60)

        weekday = weekday_name(mabel_start or bk_start or datetime.fromisoformat(day).replace(tzinfo=ET))
        row = [
            day,
            weekday,
            fmt_time(mabel_start),
            fmt_time(mabel_complete),
            round(mabel_dur / 60, 2) if mabel_dur is not None else "",
            fmt_time(bk_start),
            fmt_time(bk_complete),
            round(bk_dur / 60, 2) if bk_dur is not None else "",
            round(wait_min, 1) if wait_min is not None else "",
            round(total_min, 1) if total_min is not None else "",
            len(bk_by_day.get(day, [])),
            delay_label,
            "; ".join(notes),
        ]
        ws_pipe.append(row)
        if export and mabel_complete:
            pipeline_rows.append(
                {
                    "wait_min": wait_min,
                    "total_min": total_min,
                    "bk_complete_sec": seconds_since_midnight(bk_complete),
                    "mabel_complete_sec": seconds_since_midnight(mabel_complete),
                }
            )

    autosize_columns(ws_pipe)

    ws_sum = wb.create_sheet("summary")
    ws_sum.append(["Mabel verified -> bankruptcy matching pipeline timing"])
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum.append([])
    ws_sum.append(["analysis_period", since_label])
    ws_sum.append(["upstream_job", MABEL_JOB])
    ws_sum.append(["downstream_job", BK_JOB])
    ws_sum.append(["bk_export_detection", f"Longest succeeded run each day with duration >= {BK_MIN_EXPORT_SECONDS}s"])
    ws_sum.append([])
    ws_sum.append(["metric", "value", "notes"])
    style_header_row(ws_sum, row=7)

    summary_rows = summarize_times(mabel_runs, "mabel")
    summary_rows.append(("", "", ""))
    summary_rows.extend(summarize_times(bk_runs, "bk_all_runs"))

    if pipeline_rows:
        waits = [row["wait_min"] for row in pipeline_rows if row["wait_min"] is not None]
        totals = [row["total_min"] for row in pipeline_rows if row["total_min"] is not None]
        bk_complete_secs = [row["bk_complete_sec"] for row in pipeline_rows]
        mabel_complete_secs = [row["mabel_complete_sec"] for row in pipeline_rows]
        late_bk = sum(1 for sec in bk_complete_secs if sec >= 10 * 3600)
        very_late_bk = sum(1 for sec in bk_complete_secs if sec >= 12 * 3600)
        summary_rows.extend(
            [
                ("", "", ""),
                ("pipeline_days_with_export", str(len(pipeline_rows)), "Days with both mabel and detected BK export"),
                ("pipeline_avg_wait_mabel_to_bk_start_min", str(round(statistics.mean(waits), 1)), "After mabel completes, avg wait until BK export starts"),
                ("pipeline_median_wait_mabel_to_bk_start_min", str(round(statistics.median(waits), 1)), "Typical polling delay (BK runs every ~15 min)"),
                ("pipeline_avg_total_min", str(round(statistics.mean(totals), 1)), "Mabel start -> BK export complete"),
                ("pipeline_median_total_min", str(round(statistics.median(totals), 1)), "Median end-to-end time"),
                ("pipeline_avg_bk_matching_complete_et", seconds_to_time_str(statistics.mean(bk_complete_secs)), "When unmatched/auto_match files are ready for review"),
                ("pipeline_median_bk_matching_complete_et", seconds_to_time_str(statistics.median(bk_complete_secs)), "Median time review files land"),
                ("pipeline_earliest_bk_matching_complete_et", seconds_to_time_str(min(bk_complete_secs)), "Best case"),
                ("pipeline_latest_bk_matching_complete_et", seconds_to_time_str(max(bk_complete_secs)), "Worst case"),
                ("pipeline_days_bk_complete_after_1000_et", str(late_bk), "Days export finished at/after 10:00 ET"),
                ("pipeline_days_bk_complete_after_1200_et", str(very_late_bk), "Days export finished at/after 12:00 ET"),
                ("pipeline_avg_mabel_complete_et", seconds_to_time_str(statistics.mean(mabel_complete_secs)), "Upstream bottleneck: when verified data is ready"),
                ("pipeline_median_mabel_complete_et", seconds_to_time_str(statistics.median(mabel_complete_secs)), "Typical mabel completion"),
            ]
        )

    for row in summary_rows:
        ws_sum.append(list(row))
    autosize_columns(ws_sum)

    ws_month = wb.create_sheet("pipeline_by_month")
    ws_month.append(
        [
            "month",
            "days_with_export",
            "avg_mabel_complete_et",
            "median_mabel_complete_et",
            "avg_bk_matching_complete_et",
            "median_bk_matching_complete_et",
            "avg_total_pipeline_min",
            "median_total_pipeline_min",
        ]
    )
    style_header_row(ws_month)
    by_month: dict[str, list[dict]] = {}
    for row in ws_pipe.iter_rows(min_row=2, values_only=True):
        if not row[6]:
            continue
        month = row[0][:7]
        by_month.setdefault(month, []).append(row)

    for month in sorted(by_month):
        rows = by_month[month]
        mabel_secs = [seconds_since_midnight(datetime.strptime(f"{r[0]} {r[3]}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=ET)) for r in rows if r[3]]
        bk_secs = [seconds_since_midnight(datetime.strptime(f"{r[0]} {r[6]}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=ET)) for r in rows if r[6]]
        totals = [float(r[9]) for r in rows if r[9] not in ("", None)]
        ws_month.append(
            [
                month,
                len(rows),
                seconds_to_time_str(statistics.mean(mabel_secs)) if mabel_secs else "",
                seconds_to_time_str(statistics.median(mabel_secs)) if mabel_secs else "",
                seconds_to_time_str(statistics.mean(bk_secs)) if bk_secs else "",
                seconds_to_time_str(statistics.median(bk_secs)) if bk_secs else "",
                round(statistics.mean(totals), 1) if totals else "",
                round(statistics.median(totals), 1) if totals else "",
            ]
        )
    autosize_columns(ws_month)

    ws_notes = wb.create_sheet("dev_brief")
    ws_notes.append(["Brief: run us_mabel_verified_v001 earlier to unblock bankruptcy matching"])
    ws_notes["A1"].font = Font(bold=True, size=14)
    ws_notes.append([])

    if pipeline_rows:
        avg_mabel_complete = seconds_to_time_str(statistics.mean([row["mabel_complete_sec"] for row in pipeline_rows]))
        med_bk_complete = seconds_to_time_str(statistics.median([row["bk_complete_sec"] for row in pipeline_rows]))
        avg_bk_complete = seconds_to_time_str(statistics.mean([row["bk_complete_sec"] for row in pipeline_rows]))
        med_total = round(statistics.median([row["total_min"] for row in pipeline_rows]), 1)
        med_wait = round(statistics.median([row["wait_min"] for row in pipeline_rows]), 1)
        late_bk = sum(1 for row in pipeline_rows if row["bk_complete_sec"] >= 10 * 3600)
        points = [
            "Business impact: bankruptcy unmatched/auto_match CSVs cannot be reviewed until both jobs finish. Later mabel runs directly push back when new bankruptcies can be matched and displayed.",
            f"On {len(pipeline_rows)} days with a detected export, mabel completes on average at {avg_mabel_complete} ET.",
            f"Bankruptcy matching files (unmatched + auto_matches) land at median {med_bk_complete} ET and average {avg_bk_complete} ET.",
            f"End-to-end from mabel start to BK files ready: median {med_total} minutes (includes ~{med_wait} min polling wait after mabel finishes).",
            f"{late_bk} of {len(pipeline_rows)} days ({round(100*late_bk/len(pipeline_rows),1)}%) did not have matching files ready until 10:00 ET or later.",
            "BK job polls every ~15 minutes and skips until verified partition/data is ready, so the schedule anchor is mabel completion time.",
            "Ask: move us_mabel_verified_v001 earlier on weekdays so BK matching can finish in the morning US Eastern time.",
            "See daily_pipeline for per-day mabel complete vs BK matching complete and delay vs 08:00 ET target.",
        ]
    else:
        points = ["No pipeline rows with both mabel and BK export detected in the selected period."]

    for point in points:
        ws_notes.append([point])
    ws_notes.column_dimensions["A"].width = 120

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profile", default="<AWS_PROFILE>")
    parser.add_argument("--region", default="eu-west-1")
    parser.add_argument("--since", default="2026-05-01", help="UTC date YYYY-MM-DD for mabel history")
    parser.add_argument(
        "--bk-since",
        default="2026-06-01",
        help="UTC date YYYY-MM-DD for bankruptcy-unmatched-daily (job created 2026-06-01)",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output xlsx path (default: work/YYYY-MM-DD/mabel_bk_pipeline_timing.xlsx)",
    )
    args = parser.parse_args()

    since = datetime.fromisoformat(args.since).replace(tzinfo=UTC)
    bk_since = datetime.fromisoformat(args.bk_since).replace(tzinfo=UTC)
    today = datetime.now(ET).date()
    default_out = (
        Path(__file__).resolve().parents[1]
        / "work"
        / today.isoformat()
        / "mabel_bk_pipeline_timing.xlsx"
    )
    output = Path(args.output) if args.output else default_out

    mabel_runs = fetch_job_runs(MABEL_JOB, args.profile, args.region, since)
    bk_runs = fetch_job_runs(BK_JOB, args.profile, args.region, bk_since)
    since_label = (
        f"mabel {args.since}+, bk {args.bk_since}+ through {today.isoformat()} (ET reporting)"
    )
    build_workbook(mabel_runs, bk_runs, since_label, output)
    print(
        f"Wrote {output}\n"
        f"  mabel runs: {len(mabel_runs)}\n"
        f"  bk runs: {len(bk_runs)}"
    )


if __name__ == "__main__":
    main()
