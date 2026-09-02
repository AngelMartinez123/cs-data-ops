"""Build DET lien delete import workbook from delete-candidate IDs."""
from __future__ import annotations

import argparse
import os
from datetime import date
from pathlib import Path

import oracledb
import pandas as pd
from openpyxl import load_workbook
from shutil import copy2

CLIENT_DIR = r"C:\Users\<USERNAME>\Documents\Codex\oracle-instant-client\instantclient_23_0"
DEFAULT_DSN = "<ORACLE_HOST>:1521/<ORACLE_SERVICE>"

LIEN_COLS = [
    "rmsid",
    "safe_number",
    "date_filed",
    "filing_type",
    "amount",
    "other_case_number",
    "jurisdiction_court_code",
    "case_number",
    "action_type",
    "filed_by",
    "defendant_name",
    "defendant_suffix",
    "defendant_address",
    "defendant_city",
    "defendant_state",
    "defendant_zip_code",
    "release_date",
    "deleted_by_import",
    "deleted_by_editing_tool",
]

FETCH_SQL = """
select
    coalesce(r.rmsid, l.row_unique_id) as rmsid,
    l.safe_number,
    l.date_filed,
    l.filing_type,
    l.amount,
    l.other_case_number,
    l.jurisdiction_court_code,
    l.case_number,
    l.action_type,
    l.filed_by,
    l.defendant_name,
    l.defendant_suffix,
    l.defendant_address,
    l.defendant_city,
    l.defendant_state,
    l.defendant_zip_code,
    l.release_date
from core_data.lien l
left join source_legal_data.lien_data_rollup r on r.id = l.id
where l.id in ({placeholders})
"""


def connect():
    oracledb.init_oracle_client(lib_dir=CLIENT_DIR)
    return oracledb.connect(
        user=os.environ["ORACLE_USER"],
        password=os.environ["ORACLE_PASSWORD"],
        dsn=os.environ.get("ORACLE_DSN", DEFAULT_DSN),
    )


def load_ids(path: Path, limit: int | None) -> list[int]:
    if path.suffix.lower() == ".csv":
        col = pd.read_csv(path, nrows=0).columns[0]
        series = pd.read_csv(path, usecols=[col], dtype={col: "Int64"})[col]
    else:
        series = pd.read_excel(path, sheet_name=0, usecols=["DELETE_CANDIDATE_LIEN_ID"])[
            "DELETE_CANDIDATE_LIEN_ID"
        ]
    ids = (
        pd.to_numeric(series, errors="coerce")
        .dropna()
        .astype(int)
        .drop_duplicates()
        .tolist()
    )
    if limit:
        ids = ids[:limit]
    return ids


def fetch_rows(conn, ids: list[int], chunk: int = 900) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for i in range(0, len(ids), chunk):
        batch = ids[i : i + chunk]
        placeholders = ",".join(str(x) for x in batch)
        sql = FETCH_SQL.format(placeholders=placeholders)
        part = pd.read_sql(sql, conn)
        frames.append(part)
    if not frames:
        return pd.DataFrame(columns=LIEN_COLS[:-2])
    out = pd.concat(frames, ignore_index=True)
    out.columns = [c.lower() for c in out.columns]
    return out


def write_workbook(template: Path, output: Path, rows: pd.DataFrame, delete_date: date) -> None:
    copy2(template, output)
    wb = load_workbook(output)
    ws = wb["Lien"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    delete_val = delete_date
    for row in rows.itertuples(index=False):
        ws.append(
            [
                row.rmsid,
                row.safe_number,
                row.date_filed,
                row.filing_type,
                row.amount,
                row.other_case_number,
                row.jurisdiction_court_code,
                row.case_number,
                row.action_type,
                row.filed_by,
                row.defendant_name,
                row.defendant_suffix,
                row.defendant_address,
                row.defendant_city,
                row.defendant_state,
                row.defendant_zip_code,
                row.release_date,
                delete_val,
                delete_val,
            ]
        )
    wb.save(output)


def write_batches(
    template: Path,
    batch_dir: Path,
    ids: list[int],
    delete_date: date,
    batch_size: int,
    chunk: int,
    conn,
) -> pd.DataFrame:
    batch_dir.mkdir(parents=True, exist_ok=True)
    manifest_rows: list[dict[str, object]] = []
    total_batches = (len(ids) + batch_size - 1) // batch_size

    for batch_num, start in enumerate(range(0, len(ids), batch_size), start=1):
        batch_ids = ids[start : start + batch_size]
        rows = fetch_rows(conn, batch_ids, chunk=chunk)
        out = batch_dir / f"tax_lien_delete_batch_{batch_num:03d}_of_{total_batches:03d}.xlsx"
        write_workbook(template, out, rows, delete_date)
        manifest_rows.append(
            {
                "batch_num": batch_num,
                "file": out.name,
                "path": str(out),
                "id_count": len(batch_ids),
                "row_count": len(rows),
                "first_id": batch_ids[0],
                "last_id": batch_ids[-1],
                "status": "pending",
            }
        )
        print(f"Wrote {out.name} rows={len(rows):,}", flush=True)

    manifest = pd.DataFrame(manifest_rows)
    manifest_path = batch_dir / "batch_manifest.csv"
    manifest.to_csv(manifest_path, index=False)
    print(f"Wrote manifest {manifest_path}", flush=True)
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--ids-file",
        help="CSV or xlsx with DELETE_CANDIDATE_LIEN_ID (skips --workbook when set)",
    )
    parser.add_argument(
        "--workbook",
        help="Source workbook when --ids-file is not supplied",
    )
    parser.add_argument(
        "--template",
        required=True,
        help="Approved local import template (kept outside Git)",
    )
    parser.add_argument(
        "--output",
        default="outputs/tax_lien_delete_import.xlsx",
    )
    parser.add_argument("--delete-date", default=date.today().isoformat())
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--chunk", type=int, default=900)
    parser.add_argument(
        "--batch-size",
        type=int,
        default=None,
        help="If set, write multiple workbooks with at most this many rows each (e.g. 3000).",
    )
    parser.add_argument(
        "--batch-dir",
        default="outputs/tax_lien_delete_batches",
        help="Output directory when --batch-size is used.",
    )
    args = parser.parse_args()

    if not args.ids_file and not args.workbook:
        parser.error("provide either --ids-file or --workbook")

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    delete_date = date.fromisoformat(args.delete_date)

    source = Path(args.ids_file) if args.ids_file else Path(args.workbook)
    print(f"Loading delete-candidate IDs from {source} ...", flush=True)
    ids = load_ids(source, args.limit)
    print(f"IDs to fetch: {len(ids):,}", flush=True)

    with connect() as conn:
        if args.batch_size:
            write_batches(
                Path(args.template),
                Path(args.batch_dir),
                ids,
                delete_date,
                args.batch_size,
                args.chunk,
                conn,
            )
            return 0

        rows = fetch_rows(conn, ids, chunk=args.chunk)
    print(f"Rows fetched: {len(rows):,}", flush=True)

    missing = len(ids) - len(rows)
    if missing:
        print(f"WARNING: {missing:,} IDs not found in core_data.lien", flush=True)

    write_workbook(Path(args.template), out, rows, delete_date)
    print(f"Wrote {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
